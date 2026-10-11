"""
Text Extraction Layer (v4 — reading-order engine, tiered fetch, explicit block status)
======================================================================================

Root causes this version removes (see services/html_reader.py for the DOM side)
-------------------------------------------------------------------------------
* **Out-of-order text.** Blocks recovered by a second extractor, JSON-LD and
  hydration payloads used to be appended *after* the article's conclusion.
  Now every HTML block carries its DOM position and recovered blocks are
  inserted where they actually live on the page.
* **Partial extraction.** Text directly inside ``<div>``/``<br>`` layouts was
  never visited; ``hidden`` accordion/tab panels, FAQ ``<button>`` questions,
  the article ``<header>`` (H1) and whole-page ASP.NET ``<form>`` wrappers
  were deleted as "noise"; and trafilatura's *precision* mode could stop the
  pipeline early with sections missing.
* **Analysing challenge pages as content.** A Cloudflare "Just a moment…"
  page or a WAF 403 used to flow into the AI as if it were the article. Every
  URL now ends in an explicit status: ``ok`` · ``partial`` · ``blocked`` ·
  ``failed``. Blocked pages raise :class:`BlockedPageError` and are excluded
  from the analysis with a clear "محجوب" label.
* **Arabic mojibake.** ``charset_normalizer`` ran *before* a strict UTF-8
  attempt and could pick a single-byte codepage for valid UTF-8 Arabic. UTF-8
  validity is now checked first; double-encoded text (``Ø§Ù„``) is repaired.

Tiers (a healthy page stops after Tier 1 — one HTTP request)
------------------------------------------------------------
* Tier 0  in-process URL cache (TTL + character budget)
* Tier 1  static fetch → reading-order DOM extraction (+ trafilatura, when
          installed, used only to *recover* missed blocks in DOM order)
* Tier 2  no-network rescue: Next/Nuxt/Salla/Zid hydration, JSON-LD, then the
          site's own alternates (WordPress REST API, AMP page)
* Tier 3  real Chromium render (Playwright) for client-rendered pages and
          weak reads: lazy-load scrolling, safe accordion/tab/read-more
          expansion with a navigation lock, open shadow-DOM flattening
* Tier 4  optional reader fallback (never used to route around a detected
          block)

This module does not attempt to defeat bot protection (no CAPTCHA solving,
no fingerprint spoofing, no challenge solvers). A page that stays protected
is reported as blocked instead of being analysed on partial text.

Output contract (consumed by services/serp.py, routes and the AI prompts)
-------------------------------------------------------------------------
Blocks separated by a blank line::

    [Title]: ...
    [Description]: ...

    [H1]: main heading
    body paragraph ...

    [H2]: section heading
    • list item

    [جدول]:
    cell | cell | cell
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import threading
import time
import zipfile
from dataclasses import asdict, dataclass, field
from typing import Any, Iterable
from urllib.parse import urljoin, urlsplit

import docx
import openpyxl
import requests
from bs4 import BeautifulSoup

import config
from services import html_reader
from services.html_reader import (
    block_key,
    blocks_to_text,
    clean_line as _clean_line,
    extract_hydration_blocks,
    extract_schema_blocks,
    is_boilerplate as _is_boilerplate,
    make_soup as _soup,
    merge_in_order,
    read_html,
)
from utils.security import validate_public_http_url

try:  # Optional accelerator; the DOM engine works without it.
    import trafilatura  # type: ignore
except Exception:  # pragma: no cover - depends on the deployment
    trafilatura = None  # type: ignore

logger = logging.getLogger(config.LOGGER_NAME)

MAX_CHARS = config.MAX_CHARS

_LABELS = {
    "ar": {"heading": "عنوان", "sheet": "ورقة", "table": "جدول", "faq": "سؤال",
           "answer": "إجابة", "page": "صفحة", "caption": "شرح"},
    "en": {"heading": "Heading", "sheet": "Sheet", "table": "Table", "faq": "Question",
           "answer": "Answer", "page": "Page", "caption": "Caption"},
}


def _label(language: str, key: str) -> str:
    return _LABELS.get("en" if language == "en" else "ar", _LABELS["ar"])[key]


def _msg(language: str, ar: str, en: str) -> str:
    return en if language == "en" else ar


_WS = re.compile(r"[ \t\u00a0\u200b-\u200f\ufeff]+")
_MULTI_NL = re.compile(r"\n{3,}")
_HEADING_MARKER = re.compile(r"^\[H([1-6])\]\s*:", re.I)


# ─────────────────────────────────────────────────────────────────────────────
# Result model & errors
# ─────────────────────────────────────────────────────────────────────────────

STATUS_OK = "ok"
STATUS_PARTIAL = "partial"
STATUS_BLOCKED = "blocked"
STATUS_FAILED = "failed"


class BlockedPageError(ValueError):
    """The site actively refused automated reading (WAF / bot challenge)."""

    def __init__(self, message: str, *, provider: str = "", http_status: int | None = None, url: str = ""):
        super().__init__(message)
        self.provider = provider
        self.http_status = http_status
        self.url = url


@dataclass
class ExtractionResult:
    url: str
    text: str = ""
    status: str = STATUS_FAILED
    tier: str = ""
    http_status: int | None = None
    provider: str = ""
    reason: str = ""
    chars: int = 0
    headings: int = 0
    notes: list[str] = field(default_factory=list)

    def to_dict(self, include_text: bool = False) -> dict:
        data = asdict(self)
        if not include_text:
            data.pop("text", None)
        return data


class _HTTPStatusError(requests.HTTPError):
    def __init__(self, status: int, headers: dict | None = None, body: bytes = b""):
        super().__init__(f"HTTP {status}")
        self.status = status
        self.headers = dict(headers or {})
        self.body = body or b""


# ─────────────────────────────────────────────────────────────────────────────
# In-process URL cache (keeps ExtractionResult objects; never caches blocks)
# ─────────────────────────────────────────────────────────────────────────────

_url_memory_cache: dict[str, tuple[float, Any]] = {}
_url_memory_cache_chars: int = 0
_url_cache_lock = threading.Lock()
_browser_semaphore = threading.Semaphore(max(1, config.BROWSER_MAX_CONCURRENT))


def _cached_len(value: Any) -> int:
    return len(value.text) if isinstance(value, ExtractionResult) else len(str(value or ""))


def _drop_locked(key: str) -> None:
    global _url_memory_cache_chars
    stale = _url_memory_cache.pop(key, None)
    if stale:
        _url_memory_cache_chars = max(0, _url_memory_cache_chars - _cached_len(stale[1]))


def _cache_get(key: str) -> ExtractionResult | None:
    if not config.ENABLE_URL_MEMORY_CACHE:
        return None
    now = time.time()
    with _url_cache_lock:
        entry = _url_memory_cache.get(key)
        if not entry:
            return None
        created, value = entry
        if now - created > config.URL_MEMORY_CACHE_TTL_SECONDS:
            _drop_locked(key)
            return None
        return value if isinstance(value, ExtractionResult) else None


def _cache_set(key: str, value: ExtractionResult) -> None:
    global _url_memory_cache_chars
    if not config.ENABLE_URL_MEMORY_CACHE or not value or not value.text:
        return
    now = time.time()
    size = _cached_len(value)
    with _url_cache_lock:
        _drop_locked(key)
        for stale_key in [k for k, (created, _v) in _url_memory_cache.items()
                          if now - created > config.URL_MEMORY_CACHE_TTL_SECONDS]:
            _drop_locked(stale_key)
        while (_url_memory_cache_chars + size) > config.URL_MEMORY_CACHE_MAX_CHARS and _url_memory_cache:
            oldest = min(_url_memory_cache, key=lambda k: _url_memory_cache[k][0])
            _drop_locked(oldest)
        _url_memory_cache[key] = (now, value)
        _url_memory_cache_chars += size


# ─────────────────────────────────────────────────────────────────────────────
# HTTP
# ─────────────────────────────────────────────────────────────────────────────

_BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)

_BASE_HEADERS = {
    "User-Agent": _BROWSER_UA,
    "Accept": ("text/html,application/xhtml+xml,application/xml;q=0.9,"
               "image/avif,image/webp,*/*;q=0.8"),
    "Accept-Language": "ar,ar-SA;q=0.9,en-US;q=0.8,en;q=0.7",
    "Cache-Control": "no-cache",
    "Pragma": "no-cache",
    "Upgrade-Insecure-Requests": "1",
}

_session_local = threading.local()


def _session() -> requests.Session:
    """One pooled session per worker thread (connection reuse == speed)."""
    session = getattr(_session_local, "session", None)
    if session is None:
        session = requests.Session()
        adapter = requests.adapters.HTTPAdapter(pool_connections=16, pool_maxsize=16, max_retries=0)
        session.mount("https://", adapter)
        session.mount("http://", adapter)
        session.headers.update(_BASE_HEADERS)
        _session_local.session = session
    return session


def _read_capped(response: requests.Response, limit: int | None = None) -> bytes:
    """Stream the body but never buffer more than the configured limit."""
    limit = limit or config.MAX_URL_RESPONSE_BYTES
    chunks: list[bytes] = []
    total = 0
    for chunk in response.iter_content(chunk_size=65_536):
        if not chunk:
            continue
        chunks.append(chunk)
        total += len(chunk)
        if total >= limit:
            break
    return b"".join(chunks)[:limit]


def _fetch_public(session: requests.Session, url: str, headers: dict | None = None,
                  language: str = "ar") -> tuple[bytes, requests.Response]:
    """Fetch a public URL, validating every redirect hop against SSRF rules.

    Non-2xx answers raise :class:`_HTTPStatusError` carrying the status, the
    headers and the first 64 KB of the body so a WAF challenge can be told
    apart from an ordinary error.
    """
    current = url
    for _hop in range(max(1, config.MAX_REDIRECTS) + 1):
        response = session.get(current, timeout=(8, 20), headers=headers or None,
                               allow_redirects=False, stream=True)
        if response.status_code in (301, 302, 303, 307, 308):
            location = response.headers.get("Location", "")
            response.close()
            if not location:
                break
            current = validate_public_http_url(urljoin(current, location), language)
            continue
        if response.status_code >= 400:
            status = response.status_code
            try:
                body = _read_capped(response, 65_536)
            except Exception:
                body = b""
            response_headers = dict(response.headers or {})
            response.close()
            raise _HTTPStatusError(status, response_headers, body)
        raw = _read_capped(response)
        response.close()
        return raw, response
    raise requests.TooManyRedirects(_msg(language, "عدد التحويلات تجاوز الحد المسموح.", "Too many redirects."))


_META_CHARSET = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([a-zA-Z0-9_\-]+)""", re.I)
_MOJIBAKE = re.compile(r"[ØÙÚÛ][\u0080-\u00bf]")
_ARABIC = re.compile(r"[\u0600-\u06FF]")


def _to_single_bytes(text: str) -> bytes | None:
    """Reverse a Latin-1/CP1252 decode byte-for-byte (browsers map the five
    undefined CP1252 bytes to U+0081…U+009D, so both tables are needed)."""
    out = bytearray()
    for char in text:
        code = ord(char)
        if code < 256:
            out.append(code)
            continue
        try:
            out += char.encode("cp1252")
        except UnicodeEncodeError:
            return None
    return bytes(out)


def _repair_mojibake(text: str) -> str:
    """Undo UTF-8 that was decoded as Latin-1/CP1252 somewhere upstream
    (``Ø§Ù„`` → ``ال``) — e.g. a server that omits the charset, rendered by
    a browser that then assumes windows-1252. Applied only when the pattern
    is unmistakable and the repair yields more Arabic than it destroys."""
    if len(_MOJIBAKE.findall(text[:20000])) < 8:
        return text
    raw = _to_single_bytes(text)
    if raw is None:
        return text
    repaired = raw.decode("utf-8", errors="replace")
    if repaired.count("\ufffd") > max(3, len(_ARABIC.findall(repaired)) // 200):
        return text
    if len(_ARABIC.findall(repaired)) > len(_ARABIC.findall(text)):
        return repaired
    return text


def _decode(raw: bytes, response: Any) -> str:
    """Decode bytes with a ladder that is correct for Arabic pages.

    1. BOM   2. strict UTF-8 (valid UTF-8 is almost never accidental)
    3. declared charsets (header, <meta>)   4. Arabic code pages / detection,
    scored by Arabic-letter yield and replacement characters.
    """
    for bom, encoding in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if raw.startswith(bom):
            try:
                return _repair_mojibake(raw.decode(encoding))
            except UnicodeDecodeError:
                break
    try:
        return _repair_mojibake(raw.decode("utf-8", errors="strict"))
    except UnicodeDecodeError:
        pass

    candidates: list[str] = []
    try:
        content_type = str((getattr(response, "headers", None) or {}).get("Content-Type", ""))
    except Exception:
        content_type = ""
    match = re.search(r"charset\s*=\s*[\"']?([\w\-]+)", content_type, re.I)
    if match:
        candidates.append(match.group(1))
    meta = _META_CHARSET.search(raw[:8192])
    if meta:
        candidates.append(meta.group(1).decode("ascii", "ignore"))
    candidates += ["cp1256", "iso-8859-6"]
    try:
        import charset_normalizer
        detected = charset_normalizer.detect(raw)
        if detected and detected.get("encoding"):
            candidates.append(detected["encoding"])
    except Exception:
        pass
    candidates.append("latin-1")

    best: tuple[float, str] | None = None
    seen: set[str] = set()
    for encoding in candidates:
        key = (encoding or "").strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        try:
            text = raw.decode(encoding, errors="replace")
        except (LookupError, TypeError):
            continue
        score = len(_ARABIC.findall(text[:200_000])) - 25 * text.count("\ufffd")
        if best is None or score > best[0]:
            best = (score, text)
    return _repair_mojibake(best[1]) if best else raw.decode("utf-8", errors="ignore")


# ─────────────────────────────────────────────────────────────────────────────
# Block (WAF / bot-challenge) detection
# ─────────────────────────────────────────────────────────────────────────────

_WAF_BODY_SIGNATURES: tuple[tuple[str, re.Pattern], ...] = (
    ("Cloudflare", re.compile(
        r"cf-browser-verification|/cdn-cgi/challenge-platform|cf_chl_|cf-turnstile|"
        r"<title>\s*Just a moment\.\.\.|Attention Required! \| Cloudflare|"
        r"Checking (?:if the site connection is secure|your browser before)|"
        r"Enable JavaScript and cookies to continue", re.I)),
    ("Akamai", re.compile(r"errors\.edgesuite\.net|Reference\s*&#32;#\d|Reference\s*#\d+\.[0-9a-f]{6,}", re.I)),
    ("Imperva/Incapsula", re.compile(r"_Incapsula_Resource|Incapsula incident ID|incident_id=", re.I)),
    ("Sucuri", re.compile(r"Sucuri WebSite Firewall|Access Denied - Sucuri", re.I)),
    ("DataDome", re.compile(r"captcha-delivery\.com|geo\.captcha-delivery|datadome", re.I)),
    ("PerimeterX/HUMAN", re.compile(r"px-captcha|_pxCaptcha|perimeterx", re.I)),
    ("AWS WAF", re.compile(r"awswaf|aws-waf-token|AwsWafIntegration", re.I)),
    ("F5 BIG-IP", re.compile(r"The requested URL was rejected\. Please consult with your administrator", re.I)),
    ("Wordfence", re.compile(r"Generated by Wordfence|Your access to this site has been limited by the site owner", re.I)),
    ("ModSecurity", re.compile(r"Not Acceptable!\s*</h1>|An appropriate representation of the requested resource "
                               r"could not be found on this server\. This error was generated by Mod_Security", re.I)),
    ("CAPTCHA", re.compile(r"verify (?:that )?you are (?:a )?human|are you a robot|"
                           r"please complete the security check|تحقق من أنك (?:إنسان|انسان)", re.I)),
)


def _detect_block(status: int | None, headers: dict | None, body: str, extracted_chars: int) -> str:
    """Return the protection provider when a response is a refusal/challenge."""
    headers = {str(k).lower(): str(v) for k, v in (headers or {}).items()}
    status = status or 200
    server = headers.get("server", "").lower()
    head = (body or "")[:120_000]

    if headers.get("cf-mitigated", "").lower() == "challenge":
        return "Cloudflare"
    if "x-datadome" in headers or "x-dd-b" in headers:
        if status >= 400:
            return "DataDome"
    if "x-amzn-waf-action" in headers:
        return "AWS WAF"

    for provider, pattern in _WAF_BODY_SIGNATURES:
        if pattern.search(head):
            if status >= 400:
                return provider
            # A 200 page that merely embeds a CAPTCHA widget in a form is NOT
            # blocked — only a near-empty challenge page is.
            if extracted_chars < (2500 if provider != "CAPTCHA" else 1200):
                return provider

    if status in (401, 403, 406, 429, 451):
        if "cloudflare" in server:
            return "Cloudflare"
        if "akamai" in server:
            return "Akamai"
        if "sucuri" in server or "x-sucuri-id" in headers:
            return "Sucuri"
        if "x-iinfo" in headers or "incap" in headers.get("set-cookie", "").lower():
            return "Imperva/Incapsula"
        return f"HTTP {status}"
    if status == 503 and "cloudflare" in server:
        return "Cloudflare"
    return ""


def _blocked_message(language: str, url: str, provider: str) -> str:
    return _msg(
        language,
        f"محجوب: الموقع يمنع القراءة الآلية ({provider}) — تم استبعاده من التحليل بدل التحليل على نص ناقص: {url}",
        f"Blocked: the site refuses automated reading ({provider}) — excluded instead of analysing partial text: {url}",
    )


# ─────────────────────────────────────────────────────────────────────────────
# HTML → structured text
# ─────────────────────────────────────────────────────────────────────────────

def _trafilatura_blocks(html: str, language: str) -> list[str]:
    """trafilatura in XML/recall mode, used ONLY as a recovery oracle."""
    if trafilatura is None:
        return []
    try:
        xml = trafilatura.extract(html, output_format="xml", include_tables=True, include_formatting=False,
                                  include_comments=False, include_links=False, favor_recall=True,
                                  deduplicate=False, with_metadata=False)
    except Exception as exc:
        logger.debug("trafilatura failed: %s", exc)
        return []
    if not xml:
        return []
    try:
        tree = BeautifulSoup(xml, "xml")
    except Exception:
        tree = _soup(xml)
    blocks: list[str] = []
    for node in tree.find_all(["head", "p", "list", "table", "quote"]):
        if node.name == "list":
            blocks.extend(_clean_line(item.get_text(" ", strip=True)) for item in node.find_all("item"))
        elif node.name == "table":
            for row in node.find_all("row"):
                cells = [_clean_line(cell.get_text(" ", strip=True)) for cell in row.find_all("cell")]
                blocks.append(" | ".join(cell for cell in cells if cell))
        else:
            blocks.append(_clean_line(node.get_text(" ", strip=True)))
    return [block for block in blocks if block]


def _dedupe_blocks(blocks: Iterable[str]) -> list[str]:
    """String-level dedupe with the same rules as the DOM engine: long blocks
    are unique page-wide, short blocks/headings only lose consecutive echoes."""
    seen: set[str] = set()
    output: list[str] = []
    previous = ""
    for block in blocks:
        block = _MULTI_NL.sub("\n\n", (block or "").strip())
        if not block:
            continue
        key = block_key(block)
        if not key or key == previous:
            continue
        if not _HEADING_MARKER.match(block) and len(key) >= 40:
            if key in seen:
                continue
            seen.add(key)
        output.append(block)
        previous = key
    return output


def _merge(primary: str, extra: str) -> str:
    """Append blocks from ``extra`` that ``primary`` lacks. Used ONLY for
    content that has no DOM position (JSON-only FAQ answers, specs)."""
    if not extra:
        return primary
    if not primary:
        return extra
    existing = {block_key(block) for block in primary.split("\n\n") if block.strip()}
    additions = []
    pending_heading = ""
    for block in extra.split("\n\n"):
        block = block.strip()
        key = block_key(block)
        if not key or key in existing:
            continue
        if _HEADING_MARKER.match(block):
            pending_heading = block
            continue
        existing.add(key)
        if pending_heading:
            additions.append(pending_heading)
            pending_heading = ""
        additions.append(block)
    if not additions:
        return primary
    return primary + "\n\n" + "\n\n".join(additions)


_ARABIC_LETTER = re.compile(r"[\u0600-\u06FF\u0750-\u077F]")
_LATIN_LETTER = re.compile(r"[A-Za-z]")
_KEEP_ALWAYS = re.compile(r"^\[(?:Title|Description|H[1-6])\]\s*:", re.I)


def _script_ratio(text: str) -> float:
    arabic = len(_ARABIC_LETTER.findall(text))
    latin = len(_LATIN_LETTER.findall(text))
    total = arabic + latin
    if total < 40:
        return -1.0
    return arabic / total


def _drop_foreign_blocks(text: str) -> str:
    """Remove paragraphs written entirely in the *other* language (language
    switchers, English footers on Arabic pages). Headings, tables and blocks
    that merely contain foreign terms are never touched."""
    if not config.FILTER_FOREIGN_LANGUAGE_BLOCKS:
        return text
    blocks = [block for block in text.split("\n\n") if block.strip()]
    if len(blocks) < 3:
        return text
    ratio = _script_ratio(text)
    if ratio < 0:
        return text
    if ratio >= 0.6:
        def is_foreign(block: str) -> bool:
            return not _ARABIC_LETTER.search(block) and len(_LATIN_LETTER.findall(block)) >= 12
    elif ratio <= 0.15:
        def is_foreign(block: str) -> bool:
            return not _LATIN_LETTER.search(block) and len(_ARABIC_LETTER.findall(block)) >= 12
    else:
        return text
    kept = [block for block in blocks
            if _KEEP_ALWAYS.match(block) or block.startswith("[") or " | " in block or not is_foreign(block)]
    if len("\n\n".join(kept)) < len(text) * 0.5:
        return text
    return "\n\n".join(kept)


def _read_structured(html: str, language: str = "ar") -> tuple[str, "html_reader.ReadResult"]:
    """One HTML document → (structured text, ReadResult diagnostics)."""
    read = read_html(html, language)
    blocks = read.blocks
    traf = _trafilatura_blocks(html, language)
    if traf and blocks:
        blocks = merge_in_order(blocks, read.page_blocks, traf)
    body = blocks_to_text(blocks)

    fresh = _soup(html)
    schema = "\n\n".join(extract_schema_blocks(fresh, language))
    if schema:
        body = _merge(body, schema)
    if len(body) < config.MIN_ACCEPTABLE_EXTRACTED_CHARS:
        hydration = "\n\n".join(extract_hydration_blocks(fresh, language))
        if hydration:
            body = _merge(body, hydration)
    body = _drop_foreign_blocks(body)

    parts = [*read.meta_lines()]
    if body.strip():
        parts.append(body.strip())
    return _MULTI_NL.sub("\n\n", "\n\n".join(parts)).strip(), read


def _structured_from_html(html: str, language: str = "ar", *, base_url: str = "",
                          aggressive: bool = False) -> str:
    """Backwards-compatible entry point: HTML → structured block text."""
    if not html or len(html.strip()) < 30:
        return ""
    text, _read = _read_structured(html, language)
    return text


def _extract_main_content(html: str, language: str = "ar") -> str:
    """Backwards-compatible alias kept for older call sites and tests."""
    return _structured_from_html(html, language)


def _from_dom(soup: BeautifulSoup, language: str) -> str:
    return read_html(str(soup), language).text


def _walk_json(node: Any, out: list[str], depth: int = 0, skip_sitewide: bool = True) -> None:
    html_reader.walk_json(node, out, depth, skip_sitewide)


def _extract_embedded_content(soup: BeautifulSoup, language: str = "ar") -> str:
    return "\n\n".join(extract_hydration_blocks(soup, language))


def _extract_schema_page_content(soup: BeautifulSoup, language: str = "ar") -> str:
    return "\n\n".join(extract_schema_blocks(soup, language))


# ─────────────────────────────────────────────────────────────────────────────
# Quality gate — deterministic, so a healthy page never escalates
# ─────────────────────────────────────────────────────────────────────────────

def _count_headings(text: str) -> int:
    return len(re.findall(r"^\[H[1-6]\]\s*:", text or "", re.M))


def _body_chars(text: str) -> int:
    return sum(len(block) for block in (text or "").split("\n\n")
               if block.strip() and not re.match(r"^\[(?:Title|Description)\]\s*:", block))


def _quality(text: str) -> float:
    text = (text or "").strip()
    if not text:
        return 0.0
    blocks = [block for block in text.split("\n\n") if block.strip()]
    prose = [block for block in blocks if not _HEADING_MARKER.match(block)
             and not re.match(r"^\[(?:Title|Description)\]", block)]
    prose_chars = sum(len(block) for block in prose)
    headings = _count_headings(text)
    avg_block = prose_chars / max(len(prose), 1)
    score = float(prose_chars) + headings * 220 + min(len(prose), 120) * 18
    if avg_block < 45:
        score *= 0.55
    elif avg_block > 120:
        score *= 1.12
    boiler = sum(1 for block in prose if _is_boilerplate(block))
    score *= max(0.4, 1.0 - (boiler / max(len(prose), 1)))
    return score


def _is_strong(text: str) -> bool:
    chars = _body_chars(text)
    headings = _count_headings(text)
    if chars >= config.STRONG_EXTRACTED_CHARS and headings >= config.STRONG_EXTRACTED_HEADINGS:
        return True
    return chars >= config.STRONG_EXTRACTED_CHARS * 2 and headings >= 1


def _is_acceptable(text: str) -> bool:
    return _body_chars(text) >= config.MIN_ACCEPTABLE_EXTRACTED_CHARS


# ─────────────────────────────────────────────────────────────────────────────
# Tier 2 — the site's own alternates (no bypassing, no third parties)
# ─────────────────────────────────────────────────────────────────────────────

def _extract_content_api_text(session: requests.Session, url: str,
                              soup: BeautifulSoup | None = None, language: str = "ar") -> str:
    """WordPress REST API — only when the page is demonstrably WordPress."""
    try:
        if soup is not None:
            looks_wp = False
            generator = soup.find("meta", attrs={"name": "generator"})
            if generator and "wordpress" in str(generator.get("content") or "").lower():
                looks_wp = True
            if not looks_wp and soup.find("link", href=re.compile(r"/wp-(content|includes)/")):
                looks_wp = True
            if not looks_wp and soup.find("link", attrs={"rel": "https://api.w.org/"}):
                looks_wp = True
            if not looks_wp:
                return ""
        split = urlsplit(url)
        origin = f"{split.scheme}://{split.netloc}"
        segments = [segment for segment in split.path.split("/") if segment]
        if not segments:
            return ""
        slug = segments[-1]
        for kind in ("posts", "pages", "product"):
            endpoint = validate_public_http_url(
                f"{origin}/wp-json/wp/v2/{kind}?slug={slug}&_fields=title,content", language)
            try:
                raw, response = _fetch_public(session, endpoint, language=language)
                payload = json.loads(_decode(raw, response))
            except Exception:
                continue
            if not isinstance(payload, list) or not payload:
                continue
            item = payload[0] if isinstance(payload[0], dict) else {}
            title = (item.get("title") or {}).get("rendered") or ""
            content = (item.get("content") or {}).get("rendered") or ""
            merged = (f"<h1>{title}</h1>" if title else "") + content
            if merged.strip():
                blocks = html_reader.read_fragment(merged, language)
                return "\n\n".join(_dedupe_blocks(blocks))
    except Exception as exc:
        logger.debug("Content API fallback failed for %s: %s", url, exc)
    return ""


def _fetch_amp(session: requests.Session, url: str, soup: BeautifulSoup | None, language: str) -> str:
    """AMP alternate published by the site itself: server-rendered and lean."""
    if soup is None:
        return ""
    link = soup.find("link", attrs={"rel": re.compile(r"amphtml", re.I)})
    href = str(link.get("href") or "").strip() if link else ""
    if not href:
        return ""
    try:
        amp_url = validate_public_http_url(urljoin(url, href), language)
        raw, response = _fetch_public(session, amp_url, language=language)
        return _structured_from_html(_decode(raw, response), language)
    except Exception as exc:
        logger.debug("AMP fallback failed for %s: %s", url, exc)
        return ""


def _fetch_with_jina(url: str, language: str = "ar") -> str:
    """Optional reader service — only for weak pages that were NOT blocked."""
    if not config.ENABLE_JINA_FALLBACK:
        return ""
    try:
        response = requests.get(
            f"https://r.jina.ai/{url}", timeout=(6, 20),
            headers={"User-Agent": _BROWSER_UA, "X-Return-Format": "markdown", "X-With-Generated-Alt": "false"},
        )
        if response.status_code != 200:
            return ""
        return _markdown_to_blocks(response.text, language)
    except Exception as exc:
        logger.debug("Reader fallback failed for %s: %s", url, exc)
        return ""


# ─────────────────────────────────────────────────────────────────────────────
# Tier 3 — Chromium render (Playwright)
# ─────────────────────────────────────────────────────────────────────────────

_EXPAND_JS = r"""
() => {
  let clicked = 0;
  document.querySelectorAll('details').forEach(d => { d.open = true; });
  const selectors = [
    '[aria-expanded="false"]', '.accordion-button.collapsed', '.accordion-toggle', '.accordion-header',
    '.faq-question', '.faq__question', '[data-toggle="collapse"]', '[data-bs-toggle="collapse"]',
    '[data-bs-toggle="tab"]', '[role="tab"]', '.read-more', '.show-more', '.more-link',
    '[class*="read-more"]', '[class*="show-more"]', '[class*="readmore"]', '[class*="expand"]'
  ];
  const nav = 'nav, header, footer, [role="navigation"], [role="menubar"], [class*="menu"], [class*="navbar"], [class*="cart"]';
  document.querySelectorAll(selectors.join(',')).forEach(el => {
    if (clicked > 80) return;
    if (el.closest(nav)) return;
    const link = el.closest('a');
    if (link) {
      const href = (link.getAttribute('href') || '').trim();
      if (href && !href.startsWith('#') && !href.toLowerCase().startsWith('javascript:')) return;
    }
    try { el.click(); clicked++; } catch (e) {}
  });
  document.querySelectorAll('.collapse, .accordion-collapse, .accordion-content, .tab-pane, [role="tabpanel"], [hidden]')
    .forEach(el => {
      if (el.closest(nav)) return;
      el.classList.add('show', 'active', 'in');
      el.removeAttribute('hidden');
      el.setAttribute('aria-hidden', 'false');
      if (el.style && el.style.display === 'none') el.style.display = 'block';
    });
  return clicked;
}
"""

_SHADOW_JS = r"""
() => {
  let copied = 0;
  const walk = (root) => {
    root.querySelectorAll('*').forEach(el => {
      if (copied > 400) return;
      if (el.shadowRoot && !el.hasAttribute('data-sr-shadow')) {
        el.setAttribute('data-sr-shadow', '1');
        const holder = document.createElement('div');
        holder.setAttribute('data-sr-shadow-copy', '1');
        holder.innerHTML = el.shadowRoot.innerHTML;
        el.appendChild(holder);
        copied++;
        walk(el.shadowRoot);
      }
    });
  };
  walk(document);
  return copied;
}
"""


def _browser_request_allowed(request_url: str, verdicts: dict[str, bool]) -> bool:
    """SSRF guard for the headless browser.

    A rendered page can make its own requests (scripts, fetch/XHR, iframes).
    Without this check a malicious page could make the server's browser read
    internal addresses (cloud metadata at 169.254.169.254, localhost admin
    panels…) and the result would come back to the user inside the
    extracted text. Every http(s) request is therefore held to the same
    public-host rule as the initial URL; verdicts are cached per host:port.
    """
    lowered = (request_url or "").lower()
    if lowered.startswith(("data:", "blob:", "about:")):
        return True
    if not lowered.startswith(("http://", "https://")):
        return False
    try:
        split = urlsplit(request_url)
        key = f"{split.scheme}://{split.hostname}:{split.port or ''}"
    except ValueError:
        return False
    if key not in verdicts:
        try:
            validate_public_http_url(request_url)
            verdicts[key] = True
        except Exception:
            verdicts[key] = False
    return verdicts[key]


def _render_with_playwright(url: str, language: str = "ar", deep_read: bool = False) -> str:
    """Render with Chromium for client-rendered pages. Returns HTML ("" on failure).

    A navigation lock aborts any main-frame navigation after the first load,
    so clicking "read more"/accordion triggers can never move the browser to
    another page; the pre-expansion DOM is kept as a fallback.
    """
    if not config.ENABLE_BROWSER_FALLBACK:
        return ""
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return ""
    if not _browser_semaphore.acquire(timeout=30):
        return ""
    blocked_resources = re.compile(r"\.(png|jpe?g|gif|webp|avif|svg|ico|woff2?|ttf|otf|eot|mp4|webm|mp3|m4a)(\?|$)", re.I)
    state = {"locked": False}
    host_verdicts: dict[str, bool] = {}
    try:
        with sync_playwright() as runner:
            browser = runner.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage", "--disable-gpu"])
            try:
                context = browser.new_context(user_agent=_BROWSER_UA, locale="ar-SA",
                                              viewport={"width": 1366, "height": 2200})

                def route_handler(route):
                    request = route.request
                    try:
                        if not _browser_request_allowed(request.url, host_verdicts):
                            return route.abort()
                        if blocked_resources.search(request.url) or request.resource_type in ("image", "media", "font"):
                            return route.abort()
                        if state["locked"] and request.is_navigation_request() and request.frame == page.main_frame:
                            return route.abort()
                    except Exception:
                        return route.abort()
                    return route.continue_()

                page = context.new_page()
                page.route("**/*", route_handler)
                page.goto(url, timeout=config.BROWSER_RENDER_TIMEOUT * 1000, wait_until="domcontentloaded")
                try:
                    page.wait_for_load_state("networkidle", timeout=config.BROWSER_NETWORK_IDLE_TIMEOUT * 1000)
                except Exception:
                    pass
                state["locked"] = True
                start_url = page.url
                try:  # lazy-loaded sections
                    page.evaluate(
                        "async () => { const s = () => new Promise(r => setTimeout(r, 140));"
                        " const h = Math.min(document.body ? document.body.scrollHeight : 0, 60000);"
                        " for (let y = 0; y < h; y += 900) { window.scrollTo(0, y); await s(); }"
                        " window.scrollTo(0, 0); }")
                except Exception:
                    pass
                html_before = page.content()
                try:
                    page.evaluate(_EXPAND_JS)
                    page.wait_for_timeout(500)
                except Exception:
                    pass
                try:  # wait until the text stops growing (max ~3s)
                    previous = -1
                    for _ in range(8):
                        current = page.evaluate("() => (document.body && document.body.innerText || '').length")
                        if current == previous:
                            break
                        previous = current
                        page.wait_for_timeout(400)
                except Exception:
                    pass
                try:
                    page.evaluate(_SHADOW_JS)
                except Exception:
                    pass
                if page.url != start_url:
                    return html_before
                html = page.content()
                try:  # same-origin iframes carry the real content on some builders
                    host = urlsplit(url).netloc
                    for frame in page.frames[1:6]:
                        if urlsplit(frame.url).netloc == host:
                            html += frame.content()
                except Exception:
                    pass
                return html
            finally:
                browser.close()
    except Exception as exc:
        logger.info("Playwright render failed for %s: %s", url, exc)
        return ""
    finally:
        _browser_semaphore.release()


# ─────────────────────────────────────────────────────────────────────────────
# Markdown (reader fallback and .md uploads)
# ─────────────────────────────────────────────────────────────────────────────

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_MD_LINK = re.compile(r"\[([^\]]*)\]\([^)]*\)")


def _markdown_to_blocks(markdown: str, language: str = "ar") -> str:
    """Convert reader-style markdown into the same block contract."""
    if not markdown:
        return ""
    markdown = _MD_IMAGE.sub(" ", markdown)
    markdown = _MD_LINK.sub(r"\1", markdown)
    blocks: list[str] = []
    buffer: list[str] = []
    items: list[str] = []
    rows: list[str] = []

    def flush() -> None:
        if buffer:
            text = _clean_line(" ".join(buffer))
            if text and not _is_boilerplate(text):
                blocks.append(text)
            buffer.clear()
        if items:
            blocks.append("\n".join(items))
            items.clear()
        if rows:
            blocks.append(f"[{_label(language, 'table')}]:\n" + "\n".join(rows))
            rows.clear()

    for raw_line in markdown.replace("\r\n", "\n").split("\n"):
        line = raw_line.rstrip()
        heading = _MD_HEADING.match(line.strip())
        if heading:
            flush()
            text = _clean_line(heading.group(2))
            if text:
                blocks.append(f"[H{min(len(heading.group(1)), 6)}]: {text[:300]}")
            continue
        if not line.strip():
            flush()
            continue
        if re.match(r"^\s*(?:[-*+]|\d+[.)])\s+", line):
            if buffer or rows:
                pending = list(items)
                items.clear()
                flush()
                items.extend(pending)
            item = _clean_line(re.sub(r"^\s*(?:[-*+]|\d+[.)])\s+", "", line))
            if item and not _is_boilerplate(item):
                items.append(f"• {item}")
            continue
        if line.strip().startswith("|") and line.count("|") >= 2:
            if buffer or items:
                pending = list(rows)
                rows.clear()
                flush()
                rows.extend(pending)
            cells = [_clean_line(cell) for cell in line.strip().strip("|").split("|")]
            cells = [cell for cell in cells if cell and not re.fullmatch(r":?-{2,}:?", cell)]
            if cells:
                rows.append(" | ".join(cells))
            continue
        if items or rows:
            flush()
        buffer.append(line.strip())
    flush()
    return "\n\n".join(_dedupe_blocks(blocks))


def _truncate(text: str, limit: int = MAX_CHARS) -> str:
    """Never cut a word or a block in half."""
    text = (text or "").strip()
    if len(text) <= limit:
        return text
    window = text[:limit]
    cut = window.rfind("\n\n")
    if cut < limit * 0.6:
        cut = window.rfind(" ")
    if cut < limit * 0.5:
        cut = limit
    return window[:cut].strip()


# ─────────────────────────────────────────────────────────────────────────────
# Public API — URLs
# ─────────────────────────────────────────────────────────────────────────────

def extract_url(url: str, deep_read: bool = False, language: str = "ar") -> ExtractionResult:
    """Read a public URL and return text plus an explicit status.

    Never raises for network/content problems — those end in ``blocked`` or
    ``failed``. Invalid or non-public URLs still raise ``ValueError``.
    """
    safe_url = validate_public_http_url(url, language)
    cache_key = f"{safe_url}|{1 if deep_read else 0}"
    cached = _cache_get(cache_key)
    if cached is None and not deep_read:
        cached = _cache_get(f"{safe_url}|1")
    if cached is not None:
        return cached

    session = _session()
    result = ExtractionResult(url=safe_url)
    best, tier = "", ""
    provider = ""
    soup: BeautifulSoup | None = None
    spa_shell = False
    errors: list[str] = []

    def finish() -> ExtractionResult:
        text = _truncate(best)
        result.text = text
        result.tier = tier
        result.chars = _body_chars(text)
        result.headings = _count_headings(text)
        if provider and not _is_acceptable(text):
            result.status = STATUS_BLOCKED
            result.provider = provider
            result.reason = _blocked_message(language, safe_url, provider)
            result.text = ""
            return result
        if provider:
            result.notes.append(f"static fetch refused ({provider}); read via {tier}")
        if not text.strip():
            result.status = STATUS_FAILED
            detail = f" ({errors[0][:160]})" if errors else ""
            result.reason = _msg(language, f"تعذّر قراءة محتوى الرابط: {url}{detail}",
                                 f"Could not read content from the URL: {url}{detail}")
            return result
        result.status = STATUS_OK if _is_acceptable(text) else STATUS_PARTIAL
        if result.status == STATUS_PARTIAL:
            result.reason = _msg(language, "قراءة جزئية: النص المستخرج قصير.", "Partial read: the extracted text is short.")
        _cache_set(cache_key, result)
        return result

    # ── Tier 1 · static fetch ────────────────────────────────────────────────
    try:
        raw, response = _fetch_public(session, safe_url, language=language)
        result.http_status = int(getattr(response, "status_code", 200) or 200)
        headers = dict(getattr(response, "headers", {}) or {})
        content_type = str(headers.get("Content-Type", "")).lower()
        if "application/pdf" in content_type or safe_url.lower().endswith(".pdf"):
            best, tier = _pdf_to_text(io.BytesIO(raw), language), "static-pdf"
        elif "json" in content_type:
            collected: list[str] = []
            try:
                _walk_json(json.loads(_decode(raw, response)), collected)
            except Exception:
                pass
            best, tier = "\n\n".join(_dedupe_blocks(collected)), "static-json"
        elif content_type.startswith("text/plain"):
            lines = [line.strip() for line in _decode(raw, response).splitlines()]
            best, tier = "\n\n".join(_dedupe_blocks(_reflow_lines(lines, language))), "static-text"
        else:
            html = _decode(raw, response)
            best, read = _read_structured(html, language)
            tier = "static"
            spa_shell = read.spa_shell
            soup = _soup(html)
            provider = _detect_block(result.http_status, headers, html, _body_chars(best))
            if provider:
                best = ""
    except _HTTPStatusError as exc:
        result.http_status = exc.status
        body = _decode(exc.body, None) if exc.body else ""
        provider = _detect_block(exc.status, exc.headers, body, 0)
        errors.append(str(exc))
        logger.info("Static fetch for %s returned HTTP %s%s", safe_url, exc.status,
                    f" ({provider})" if provider else "")
    except Exception as exc:
        errors.append(str(exc))
        logger.info("Static fetch failed for %s: %s", safe_url, exc)

    if _is_strong(best) and not spa_shell:
        return finish()

    # ── Tier 2 · site-provided alternates (no network for hydration) ─────────
    if soup is not None and not provider and not _is_strong(best):
        for fetcher, name in ((lambda: _extract_content_api_text(session, safe_url, soup, language), "cms-api"),
                              (lambda: _fetch_amp(session, safe_url, soup, language), "amp")):
            if _is_acceptable(best) and not spa_shell:
                break
            try:
                candidate = fetcher()
            except Exception as exc:
                errors.append(str(exc))
                candidate = ""
            if candidate and _quality(candidate) > _quality(best) * 1.1:
                best, tier = candidate, name
    if _is_strong(best) and not spa_shell:
        return finish()

    # ── Tier 3 · real browser render ─────────────────────────────────────────
    if not _is_acceptable(best) or spa_shell or provider or (deep_read and not _is_strong(best)):
        rendered = _render_with_playwright(safe_url, language, deep_read)
        if rendered:
            rendered = _repair_mojibake(rendered)
            candidate, _read = _read_structured(rendered, language)
            rendered_provider = _detect_block(200, {}, rendered, _body_chars(candidate))
            if not rendered_provider and _quality(candidate) > _quality(best):
                best, tier = candidate, "browser"
            elif rendered_provider and not provider and not _is_acceptable(best):
                provider = rendered_provider

    # ── Tier 4 · reader fallback (never for protected pages) ────────────────
    if not provider and not _is_acceptable(best):
        reader_text = _fetch_with_jina(safe_url, language)
        if _quality(reader_text) > _quality(best):
            best, tier = reader_text, "reader"

    return finish()


def extract_text_from_url(url: str, deep_read: bool = False, language: str = "ar") -> str:
    """String API kept for existing callers. Raises :class:`BlockedPageError`
    for protected pages and ``ValueError`` when nothing could be read."""
    result = extract_url(url, deep_read, language)
    if result.status == STATUS_BLOCKED:
        raise BlockedPageError(result.reason, provider=result.provider,
                               http_status=result.http_status, url=result.url)
    if result.status == STATUS_FAILED:
        raise ValueError(result.reason or _msg(language, f"تعذّر قراءة محتوى الرابط: {url}",
                                               f"Could not read content from the URL: {url}"))
    return result.text


# ─────────────────────────────────────────────────────────────────────────────
# Public API — files
# ─────────────────────────────────────────────────────────────────────────────

def _sanitize_code(raw: str) -> str:
    """Strip code noise, keep identifiers and comments for entity extraction."""
    raw = re.sub(r'"[^"]*"', " ", raw)
    raw = re.sub(r"'[^']*'", " ", raw)
    raw = re.sub(r"https?://\S+", " ", raw)
    return re.sub(r"\s+", " ", raw).strip()


def _decode_bytes(raw: bytes) -> str:
    for bom, encoding in ((b"\xef\xbb\xbf", "utf-8-sig"),
                          (b"\xff\xfe", "utf-16"),
                          (b"\xfe\xff", "utf-16")):
        if raw.startswith(bom):
            try:
                return raw.decode(encoding)
            except Exception:
                break
    for encoding in ("utf-8", "cp1256", "iso-8859-6", "latin-1"):
        try:
            text = raw.decode(encoding, errors="strict")
            if "\ufffd" not in text:
                return text
        except (LookupError, UnicodeDecodeError):
            continue
    return raw.decode("utf-8", errors="ignore")


_TERMINAL_PUNCT = ".!?:;\u060c\u061b\u061f\u06d4\u2026)\u00bb\"'"
_PAGE_NUMBER_LINE = re.compile(
    r"^(?:page\s*)?\d{1,4}(?:\s*(?:/|of|\u0645\u0646)\s*\d{1,4})?$|"
    r"^\u0635\u0641\u062d\u0629\s*\d{1,4}(?:\s*\u0645\u0646\s*\d{1,4})?$|"
    r"^[ivxlcdm]{1,7}$",
    re.I,
)
_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+(\S.*)$")


def _strip_running_chrome(pages: list[list[str]]) -> list[list[str]]:
    """Remove the header/footer band that repeats on most pages.

    A PDF's running header ("Company | Internal Report"), its footer
    ("Confidential"), and its page numbers are the direct equivalent of a web
    page's <header>/<footer> — they are not content and they poison entity
    extraction by appearing once per page.
    """
    if len(pages) < 2:
        return [[line for line in page if not _PAGE_NUMBER_LINE.fullmatch(line.strip())]
                for page in pages]

    counts: dict[str, int] = {}
    for page in pages:
        # Chrome lives in the first and last few lines of a page.
        band = page[:4] + page[-4:]
        for line in {re.sub(r"\s+", " ", item).strip().lower() for item in band if item.strip()}:
            counts[line] = counts.get(line, 0) + 1

    threshold = max(2, int(len(pages) * 0.6))
    chrome = {line for line, hits in counts.items() if hits >= threshold and len(line) <= 160}

    cleaned: list[list[str]] = []
    for page in pages:
        kept = []
        for line in page:
            stripped = line.strip()
            if not stripped:
                continue
            if _PAGE_NUMBER_LINE.fullmatch(stripped):
                continue
            if re.sub(r"\s+", " ", stripped).lower() in chrome:
                continue
            kept.append(stripped)
        cleaned.append(kept)
    return cleaned


def _reflow_lines(lines: list[str], language: str = "ar") -> list[str]:
    """Rebuild paragraphs and headings from hard-wrapped PDF lines.

    PDF text has one entry per *visual* line, so a single paragraph arrives as
    four fragments. Joining them back is what makes sentence-level analysis
    (questions, proof signals, topic sentences) work downstream.
    """
    lengths = [len(line) for line in lines if len(line) > 20]
    wrap_width = max(lengths) if lengths else 90

    blocks: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = _clean_line(" ".join(buffer))
            if text and not _is_boilerplate(text):
                blocks.append(text)
            buffer.clear()

    def is_heading(index: int, line: str) -> tuple[bool, int]:
        if len(line) > 110 or len(line) < 3:
            return False, 2
        match = _NUMBERED_HEADING.match(line)
        if match and len(match.group(2)) <= 90:
            depth = match.group(1).count(".") + 1
            return True, min(depth + 1, 4)
        if line[-1] in _TERMINAL_PUNCT:
            return False, 2
        # A line is a heading only when it is *noticeably* shorter than the
        # body wrap width and sits directly above a full-width body line.
        # Comparing against the measured wrap width (rather than a fixed
        # character count) is what stops the first line of a hard-wrapped
        # paragraph from being mistaken for a section title.
        if len(line) >= wrap_width * 0.7:
            return False, 2
        following = lines[index + 1] if index + 1 < len(lines) else ""
        if len(following) >= wrap_width * 0.7:
            return True, 2
        return False, 2

    for index, line in enumerate(lines):
        heading, level = is_heading(index, line)
        if heading:
            flush()
            text = _clean_line(line)
            if text:
                blocks.append(f"[H{level}]: {text[:300]}")
            continue
        buffer.append(line)
        # A line noticeably shorter than the wrap width ends the paragraph.
        if line and (line[-1] in _TERMINAL_PUNCT or len(line) < wrap_width * 0.62):
            flush()
    flush()
    return blocks


def _pdf_to_text(stream: io.BytesIO, language: str = "ar") -> str:
    """Extract PDF body text: running header/footer removed, paragraphs
    reflowed, and headings promoted to the same markers URLs produce."""
    try:
        from pypdf import PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except Exception as exc:
            raise ValueError(
                _msg(language,
                     "\u0642\u0631\u0627\u0621\u0629 \u0645\u0644\u0641\u0627\u062a PDF \u062a\u062d\u062a\u0627\u062c \u062a\u062b\u0628\u064a\u062a \u0645\u0643\u062a\u0628\u0629 pypdf.",
                     "Reading PDF files requires the pypdf package.")
            ) from exc
    stream.seek(0)
    reader = PdfReader(stream)

    pages: list[list[str]] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append([line.strip() for line in text.split("\n")])

    cleaned_pages = _strip_running_chrome(pages)
    lines = [line for page in cleaned_pages for line in page if line]
    if not lines:
        return ""
    blocks = _reflow_lines(lines, language)
    # The first short unpunctuated line of a document is its title.
    if blocks and not _HEADING_MARKER.match(blocks[0]):
        first = blocks[0]
        if len(first) <= 110 and first[-1] not in _TERMINAL_PUNCT:
            blocks[0] = f"[H1]: {first[:300]}"
    return "\n\n".join(_dedupe_blocks(blocks))


def _docx_heading_level(style_name: str) -> int | None:
    """Map a Word paragraph style to an HTML heading level, or None."""
    style = (style_name or "").strip().lower()
    if style in ("title", "\u0627\u0644\u0639\u0646\u0648\u0627\u0646"):
        return 1
    if style.startswith("subtitle"):
        return 2
    match = re.match(r"^heading\s*(\d+)", style)
    if match:
        return min(int(match.group(1)), 6)
    # Localised Word installs name the styles in the UI language.
    match = re.match(r"^\u0639\u0646\u0648\u0627\u0646\s*(\d+)", style)
    if match:
        return min(int(match.group(1)), 6)
    return None


def _docx_is_list(paragraph, style_name: str) -> bool:
    if "list" in (style_name or "").lower() or "\u0642\u0627\u0626\u0645\u0629" in (style_name or ""):
        return True
    try:  # numbered/bulleted via numbering properties rather than style
        p_pr = paragraph._p.pPr
        return p_pr is not None and p_pr.numPr is not None
    except Exception:
        return False


def _docx_to_text(stream: io.BytesIO, language: str = "ar") -> str:
    """Read a Word document in true document order.

    Headings use the same ``[H1]..[H6]`` markers a URL produces, so a DOCX
    uploaded as "my site" and a competitor URL are analysed on identical
    footing. Word headers and footers live outside ``document.paragraphs``
    and are therefore excluded automatically.
    """
    stream.seek(0)
    document = docx.Document(stream)
    table_label = _label(language, "table")

    paragraphs = {id(item._p): item for item in document.paragraphs}
    tables = {id(item._tbl): item for item in document.tables}

    blocks: list[str] = []
    pending_list: list[str] = []

    def flush_list() -> None:
        if pending_list:
            blocks.append("\n".join(pending_list))
            pending_list.clear()

    for child in document.element.body.iterchildren():
        paragraph = paragraphs.get(id(child))
        if paragraph is not None:
            text = _clean_line(paragraph.text)
            if not text:
                continue
            style = getattr(paragraph.style, "name", "") or ""
            if style.lower().startswith(("toc", "table of contents", "index")):
                continue  # a table of contents is navigation, not content

            level = _docx_heading_level(style)
            if level is not None:
                flush_list()
                blocks.append(f"[H{level}]: {text[:300]}")
                continue

            if _is_boilerplate(text):
                continue

            if _docx_is_list(paragraph, style):
                pending_list.append(f"\u2022 {text}")
                continue

            flush_list()
            blocks.append(text)
            continue

        table = tables.get(id(child))
        if table is not None:
            flush_list()
            rows = []
            for row in table.rows:
                cells = [_clean_line(cell.text) for cell in row.cells]
                cells = [cell for cell in cells if cell]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(f"[{table_label}]:\n" + "\n".join(rows))
    flush_list()

    return "\n\n".join(_dedupe_blocks(blocks))


def _xlsx_to_text(stream: io.BytesIO, language: str = "ar") -> str:
    stream.seek(0)
    workbook = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    sheet_label = _label(language, "sheet")
    blocks: list[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row
                         if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(f"[{sheet_label}]: {sheet_name}")
                blocks.append("\n".join(rows))
    finally:
        try:
            workbook.close()
        except Exception:
            pass
    return "\n\n".join(blocks)


def _pptx_to_text(stream: io.BytesIO, language: str = "ar") -> str:
    """Read slide text straight from the OOXML package (no extra dependency)."""
    stream.seek(0)
    blocks: list[str] = []
    page_label = _msg(language, "\u0634\u0631\u064a\u062d\u0629", "Slide")
    with zipfile.ZipFile(stream) as archive:
        slides = sorted(
            (name for name in archive.namelist()
             if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.findall(r"\d+", name)[-1]),
        )
        for index, name in enumerate(slides, start=1):
            xml = archive.read(name).decode("utf-8", errors="ignore")
            texts = [_clean_line(chunk) for chunk in re.findall(r"<a:t[^>]*>(.*?)</a:t>", xml, re.S)]
            texts = [text for text in texts if text]
            if texts:
                blocks.append(f"[{page_label} {index}]")
                blocks.append("\n".join(texts))
    return "\n\n".join(_dedupe_blocks(blocks))


def _csv_to_text(text: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        cells = [str(cell).strip() for cell in row if str(cell).strip()]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_text_from_file(file_obj, deep_read: bool = False, language: str = "ar") -> str:
    """Read an uploaded file and return the same structured block contract."""
    filename = str(getattr(file_obj, "filename", "unknown") or "unknown").lower()

    try:
        try:
            file_obj.seek(0)
        except Exception:
            pass
        raw = file_obj.read()
        if isinstance(raw, str):
            raw = raw.encode("utf-8", errors="ignore")
        raw = raw[:config.MAX_FILE_BYTES]
        stream = io.BytesIO(raw)

        if filename.endswith((".doc", ".xls", ".ppt")):
            raise ValueError(
                _msg(language,
                     "صيغة الملف القديمة غير مدعومة. برجاء الحفظ بصيغة docx أو xlsx أو pptx.",
                     "This legacy file format is not supported. Please save as docx, xlsx, or pptx.")
            )

        if filename.endswith(".pdf"):
            return _truncate(_pdf_to_text(stream, language))
        if filename.endswith(".docx"):
            return _truncate(_docx_to_text(stream, language))
        if filename.endswith((".xlsx", ".xlsm", ".xltx")):
            return _truncate(_xlsx_to_text(stream, language))
        if filename.endswith(".pptx"):
            return _truncate(_pptx_to_text(stream, language))
        if filename.endswith((".html", ".htm", ".xhtml")):
            return _truncate(_structured_from_html(_decode_bytes(raw), language, aggressive=True))
        if filename.endswith((".csv", ".tsv")):
            return _truncate(_csv_to_text(_decode_bytes(raw)))
        if filename.endswith((".md", ".markdown", ".mdx")):
            return _truncate(_markdown_to_blocks(_decode_bytes(raw), language))
        if filename.endswith(".json"):
            collected: list[str] = []
            try:
                _walk_json(json.loads(_decode_bytes(raw)), collected)
            except Exception:
                return _truncate(_decode_bytes(raw))
            return _truncate("\n\n".join(_dedupe_blocks(collected)))
        if filename.endswith((".xml", ".rss", ".atom")):
            try:
                from defusedxml.ElementTree import fromstring
                root = fromstring(_decode_bytes(raw))
                texts = [_clean_line(node.text or "") for node in root.iter()]
                return _truncate("\n\n".join(_dedupe_blocks(text for text in texts if text)))
            except Exception:
                return _truncate(_from_dom(_soup(_decode_bytes(raw)), language))
        if filename.endswith((".py", ".js", ".ts", ".css", ".jsx", ".tsx",
                              ".java", ".cpp", ".cs", ".php", ".go", ".rb",
                              ".rs", ".swift", ".kt", ".sql", ".sh")):
            return _truncate(_sanitize_code(_decode_bytes(raw)))
        if filename.endswith((".txt", ".text", ".log")):
            # Plain text is very often hard-wrapped at a fixed column; reflow
            # it so sentences survive intact for the analysis.
            lines = [line.strip() for line in _decode_bytes(raw).splitlines()]
            return _truncate("\n\n".join(_dedupe_blocks(_reflow_lines(lines, language))))

        # Unknown extension — sniff the magic bytes before giving up.
        if raw[:4] == b"%PDF":
            return _truncate(_pdf_to_text(stream, language))
        if raw[:2] == b"PK":
            try:
                return _truncate(_docx_to_text(stream, language))
            except Exception:
                pass
        text = _decode_bytes(raw)
        if "<html" in text[:4000].lower() or "<!doctype html" in text[:200].lower():
            return _truncate(_structured_from_html(text, language, aggressive=True))
        return _truncate(_MULTI_NL.sub("\n\n", text))

    except ValueError:
        raise
    except Exception as exc:
        logger.error("File extraction failed for %s: %s", filename, exc)
        raise ValueError(
            _msg(language, f"تعذّر قراءة الملف: {filename}", f"Could not read the file: {filename}")
        ) from exc
