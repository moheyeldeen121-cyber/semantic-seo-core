"""Export canonical semantic HTML as a native, RTL-aware Word document."""

from __future__ import annotations

import io

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from services.rich_content import normalize_rich_document

_HEADING_SIZES_PT = {1: 22, 2: 17, 3: 14, 4: 12}
_NORMAL_SIZE_PT = 11


def _set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attribute}"), font_name)


def _set_style_font(style, font_name: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attribute}"), font_name)
    sz_cs = rpr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = rpr.makeelement(qn("w:szCs"), {})
        rpr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    if bold and rpr.find(qn("w:bCs")) is None:
        rpr.append(rpr.makeelement(qn("w:bCs"), {}))


def _set_bidi(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(ppr.makeelement(qn("w:bidi"), {}))


def _setup_document_styles(doc: Document, language: str) -> None:
    font_name = "Alexandria" if language == "ar" else "Calibri"
    _set_style_font(doc.styles["Normal"], font_name, _NORMAL_SIZE_PT)
    for level, size_pt in _HEADING_SIZES_PT.items():
        style_name = f"Heading {level}"
        if style_name in doc.styles:
            _set_style_font(doc.styles[style_name], font_name, size_pt, bold=True)


def _new_paragraph(doc, language: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if language == "ar" else WD_ALIGN_PARAGRAPH.LEFT
    if language == "ar":
        _set_bidi(paragraph)
    return paragraph


def _add_inline_runs(paragraph, node, *, bold: bool = False, italic: bool = False) -> None:
    if isinstance(node, NavigableString):
        value = str(node)
        if value:
            run = paragraph.add_run(value)
            run.bold = bold
            run.italic = italic
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name == "br":
        paragraph.add_run().add_break()
        return
    next_bold = bold or name in {"strong", "b", "th"}
    next_italic = italic or name in {"em", "i"}
    for child in node.children:
        _add_inline_runs(paragraph, child, bold=next_bold, italic=next_italic)


def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(tc_pr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color}))


def _add_table(doc: Document, table_node: Tag, language: str) -> None:
    source_rows = table_node.find_all("tr")
    if not source_rows:
        return
    col_count = max(
        (len(row.find_all(["th", "td"], recursive=False)) for row in source_rows),
        default=1,
    )
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    if language == "ar":
        tbl_pr = table._tbl.tblPr
        if tbl_pr.find(qn("w:bidiVisual")) is None:
            tbl_pr.append(tbl_pr.makeelement(qn("w:bidiVisual"), {}))

    for row_index, source_row in enumerate(source_rows):
        target_row = table.add_row()
        source_cells = source_row.find_all(["th", "td"], recursive=False)
        for col_index in range(col_count):
            cell = target_row.cells[col_index]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if language == "ar" else WD_ALIGN_PARAGRAPH.LEFT
            if language == "ar":
                _set_bidi(paragraph)
            if col_index < len(source_cells):
                source_cell = source_cells[col_index]
                _add_inline_runs(paragraph, source_cell, bold=source_cell.name == "th")
                is_header = source_cell.name == "th"
            else:
                is_header = False
            if row_index == 0 or is_header:
                for run in paragraph.runs:
                    run.bold = True
                _shade_cell(cell, "E7E6E6")


def _render_node(doc: Document, node, language: str, list_style: str | None = None) -> None:
    if isinstance(node, NavigableString):
        text = str(node).strip()
        if text:
            _new_paragraph(doc, language).add_run(text)
        return
    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name in {"article", "section", "body", "html"}:
        for child in node.children:
            _render_node(doc, child, language)
        return
    if name in {"h1", "h2", "h3", "h4"}:
        paragraph = _new_paragraph(doc, language, style=f"Heading {name[1]}")
        _add_inline_runs(paragraph, node, bold=True)
        return
    if name == "p":
        paragraph = _new_paragraph(doc, language)
        _add_inline_runs(paragraph, node)
        return
    if name == "blockquote":
        paragraph = _new_paragraph(doc, language, style="Quote")
        _add_inline_runs(paragraph, node)
        return
    if name in {"ul", "ol"}:
        style = "List Bullet" if name == "ul" else "List Number"
        for item in node.find_all("li", recursive=False):
            _render_node(doc, item, language, list_style=style)
        return
    if name == "li":
        paragraph = _new_paragraph(doc, language, style=list_style or "List Bullet")
        for child in node.children:
            if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                continue
            _add_inline_runs(paragraph, child)
        for nested in node.find_all(["ul", "ol"], recursive=False):
            _render_node(doc, nested, language)
        return
    if name == "table":
        _add_table(doc, node, language)
        return
    if name == "hr":
        _new_paragraph(doc, language).add_run("─" * 30)
        return

    for child in node.children:
        _render_node(doc, child, language)


def build_docx_report(content: str, language: str = "ar") -> bytes:
    """Convert semantic HTML (or legacy Markdown fallback) to native DOCX."""
    doc = Document()
    if language == "ar":
        section_properties = doc.sections[0]._sectPr
        if section_properties.find(qn("w:bidi")) is None:
            section_properties.append(section_properties.makeelement(qn("w:bidi"), {}))
    _setup_document_styles(doc, language)

    soup = BeautifulSoup(normalize_rich_document(content, language), "html.parser")
    article = soup.find("article") or soup
    for child in article.children:
        _render_node(doc, child, language)

    font_name = "Alexandria" if language == "ar" else "Calibri"
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            _set_run_font(run, run.font.name or font_name)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
