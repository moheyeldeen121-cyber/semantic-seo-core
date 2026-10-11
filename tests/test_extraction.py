import io

import docx
import openpyxl

from services.extraction import extract_text_from_file


class NamedBytesIO(io.BytesIO):
    def __init__(self, payload: bytes, filename: str):
        super().__init__(payload)
        self.filename = filename


def test_docx_is_seekable_and_extracts_headings_and_tables():
    document = docx.Document()
    document.add_heading("عنوان المقال", level=1)
    document.add_paragraph("محتوى تجريبي واضح")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "كيان"
    table.cell(0, 1).text = "قيمة"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text_from_file(NamedBytesIO(buffer.getvalue(), "article.docx"))
    assert "[H1]: عنوان المقال" in text
    assert "محتوى تجريبي واضح" in text
    assert "كيان | قيمة" in text


def test_xlsx_extracts_sheet_and_cells():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Entities"
    sheet.append(["الكيان", "النوع"])
    sheet.append(["SEO", "Concept"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = extract_text_from_file(NamedBytesIO(buffer.getvalue(), "entities.xlsx"))
    assert "[ورقة]: Entities" in text
    assert "الكيان | النوع" in text
    assert "SEO | Concept" in text


def test_legacy_xls_is_rejected_explicitly():
    try:
        extract_text_from_file(NamedBytesIO(b"not-an-xls", "legacy.xls"))
    except ValueError as exc:
        assert "غير مدعومة" in str(exc)
    else:
        raise AssertionError("Legacy .xls should not be accepted")


def test_hydration_state_extracts_salla_zid_style_content():
    from bs4 import BeautifulSoup
    from services.extraction import _extract_embedded_content

    html = '''<html><body><script id="__salla_data__">
    {"product":{"name":"منتج احترافي","description":"وصف تفصيلي للمنتج والخدمة المرتبطة به"}}
    </script><script>window.__ZID_DATA__ = {"question":"كيف يتم الشحن؟","answer":"يتم الشحن خلال يومين"};</script></body></html>'''
    text = _extract_embedded_content(BeautifulSoup(html, "html.parser"), "ar")
    assert "منتج احترافي" in text
    assert "كيف يتم الشحن؟" in text
    assert "يتم الشحن خلال يومين" in text


def test_strong_static_page_skips_duplicate_reader_and_browser(monkeypatch):
    import services.extraction as extraction

    paragraphs = "".join(f"<p>Detailed paragraph {i} about pricing, process, evidence, and service delivery.</p>" for i in range(14))
    html = f"<html><body><article><h1>Main topic</h1><h2>Section one</h2><h2>Section two</h2>{paragraphs}</article></body></html>"

    class Response:
        headers = {"Content-Type": "text/html; charset=utf-8"}
        encoding = "utf-8"

    calls = {"fetch": 0, "cms": 0, "browser": 0}

    def fake_fetch(session, url, headers=None, language="ar"):
        calls["fetch"] += 1
        return html.encode(), Response()

    def fake_cms(*args, **kwargs):
        calls["cms"] += 1
        return ""

    def fake_browser(*args, **kwargs):
        calls["browser"] += 1
        return ""

    monkeypatch.setattr(extraction, "validate_public_http_url", lambda url, language="ar": url)
    monkeypatch.setattr(extraction, "_fetch_public", fake_fetch)
    monkeypatch.setattr(extraction, "_extract_content_api_text", fake_cms)
    monkeypatch.setattr(extraction, "_render_with_playwright", fake_browser)
    monkeypatch.setattr(extraction.config, "STRONG_EXTRACTED_CHARS", 500)
    monkeypatch.setattr(extraction.config, "STRONG_EXTRACTED_HEADINGS", 3)
    monkeypatch.setattr(extraction.config, "MIN_ACCEPTABLE_EXTRACTED_CHARS", 500)
    monkeypatch.setattr(extraction.config, "ENABLE_JINA_FALLBACK", True)
    monkeypatch.setattr(extraction.config, "ENABLE_BROWSER_FALLBACK", True)

    text = extraction.extract_text_from_url("https://public.example/page", deep_read=True, language="en")
    assert "Main topic" in text
    assert calls == {"fetch": 1, "cms": 0, "browser": 0}


def test_url_memory_cache_avoids_duplicate_fetch(monkeypatch):
    import services.extraction as extraction

    paragraphs = "".join(
        f"<p>Complete paragraph {i} with enough useful semantic content for caching.</p>"
        for i in range(16)
    )
    html = (
        "<html><body><article><h1>Cached topic</h1><h2>First section</h2>"
        f"<h2>Second section</h2>{paragraphs}</article></body></html>"
    )

    class Response:
        headers = {"Content-Type": "text/html; charset=utf-8"}
        encoding = "utf-8"

    calls = {"fetch": 0}

    def fake_fetch(session, url, headers=None, language="ar"):
        calls["fetch"] += 1
        return html.encode(), Response()

    with extraction._url_cache_lock:
        extraction._url_memory_cache.clear()
        extraction._url_memory_cache_chars = 0

    monkeypatch.setattr(extraction, "validate_public_http_url", lambda url, language="ar": url)
    monkeypatch.setattr(extraction, "_fetch_public", fake_fetch)
    monkeypatch.setattr(extraction.config, "ENABLE_URL_MEMORY_CACHE", True)
    monkeypatch.setattr(extraction.config, "URL_MEMORY_CACHE_TTL_SECONDS", 600)
    monkeypatch.setattr(extraction.config, "STRONG_EXTRACTED_CHARS", 500)
    monkeypatch.setattr(extraction.config, "STRONG_EXTRACTED_HEADINGS", 3)
    monkeypatch.setattr(extraction.config, "MIN_ACCEPTABLE_EXTRACTED_CHARS", 500)
    monkeypatch.setattr(extraction.config, "ENABLE_JINA_FALLBACK", False)
    monkeypatch.setattr(extraction.config, "ENABLE_BROWSER_FALLBACK", False)

    first = extraction.extract_text_from_url(
        "https://cache-test.example/unique-page", deep_read=True, language="en"
    )
    second = extraction.extract_text_from_url(
        "https://cache-test.example/unique-page", deep_read=True, language="en"
    )

    assert first == second
    assert "Cached topic" in second
    assert calls["fetch"] == 1
