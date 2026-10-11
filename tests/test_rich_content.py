import io

import docx

from services.docx_export import build_docx_report
from services.rich_content import html_to_plain_text, normalize_rich_document
from services.serp import _outline_quality


def test_markdown_fallback_becomes_clean_semantic_html():
    html = normalize_rich_document("# عنوان\n\n## قسم\n\nنص **مهم**", "ar")
    assert "<h1>عنوان</h1>" in html
    assert "<h2>قسم</h2>" in html
    assert "#" not in html_to_plain_text(html)
    assert "**" not in html_to_plain_text(html)


def test_dangerous_html_is_removed_not_unwrapped():
    html = normalize_rich_document("<article><h1>آمن</h1><script>alert(1)</script><p>نص</p></article>", "ar")
    assert "script" not in html.lower()
    assert "alert(1)" not in html
    assert "<h1>آمن</h1>" in html


def test_docx_uses_native_heading_levels_and_clean_text():
    payload = build_docx_report(
        "<article><h1>العنوان</h1><p>مقدمة <strong>مهمة</strong></p><h2>قسم</h2><ul><li>نقطة</li></ul></article>",
        language="ar",
    )
    document = docx.Document(io.BytesIO(payload))
    texts = [p.text for p in document.paragraphs]
    styles = [p.style.name for p in document.paragraphs]
    assert "العنوان" in texts
    assert "قسم" in texts
    assert any(style.startswith("Heading 1") for style in styles)
    assert any(style.startswith("Heading 2") for style in styles)
    assert all("**" not in text and not text.startswith("#") for text in texts)


def test_outline_quality_rejects_shallow_or_orphan_structure():
    quality = _outline_quality("<article><h1>عنوان</h1><h2>قسم وحيد</h2></article>")
    assert not quality["valid"]
    assert quality["counts"]["h2"] == 1
    assert quality["orphan_headings"]
