"""End-to-end /analyze: extraction statuses → AI (mocked) → salience → clusters → Excel."""
import io
import json

import openpyxl
import pytest

pytest.importorskip("flask")

import routes.analyze as analyze
import services.claude as claude
import services.verification as verification
from services.extraction import ExtractionResult

MY = ("[Title]: دليل نية البحث\n\n[H1]: دليل نية البحث\n\nنية البحث تحدد نوع المحتوى والكلمات المفتاحية المناسبة للمستخدم.\n\n"
      "[H2]: الكلمات المفتاحية\n\nاختيار الكلمات المفتاحية يعتمد على نية البحث وحجم البحث.")
COMP = ("[H1]: دليل السيو\n\nالسيو التقني يشمل الفهرسة وخريطة الموقع وسرعة الموقع.\n\n[H2]: السيو التقني\n\n"
        "الفهرسة وخريطة الموقع أساس السيو التقني، ونية البحث مهمة أيضاً.")

ENTITIES = {
    "mine": {"main_topic": "نية البحث", "clusters": [{"cluster_name": "الكلمات", "pillar": "نية البحث", "entities": [
        {"name": "نية البحث", "source_term": "نية البحث", "type": "Concept", "importance": "high"},
        {"name": "الكلمات المفتاحية", "source_term": "الكلمات المفتاحية", "type": "Concept", "importance": "high"}]}]},
    "comp": {"main_topic": "السيو", "clusters": [{"cluster_name": "السيو التقني", "pillar": "السيو التقني", "entities": [
        {"name": "السيو التقني", "source_term": "السيو التقني", "type": "Concept", "importance": "high"},
        {"name": "الفهرسة", "source_term": "الفهرسة", "type": "Concept", "importance": "medium"},
        {"name": "خريطة الموقع", "source_term": "خريطة الموقع", "type": "Concept", "importance": "medium"}]},
        {"cluster_name": "البحث", "pillar": "نية البحث", "entities": [
        {"name": "نية البحث", "source_term": "نية البحث", "type": "Concept", "importance": "medium"}]}]},
}


def _fake_url_result(url, deep_read, language):
    if "blocked" in url:
        return ExtractionResult(url=url, status="blocked", provider="Cloudflare", reason="محجوب: Cloudflare")
    return ExtractionResult(url=url, text=MY if "mine" in url else COMP, status="ok", tier="static", chars=900, headings=2)


def _fake_ai(prompt, system_prompt, model, api_key, api_url="", max_tokens=0, task_type="", return_usage=False, **kw):
    payload = ENTITIES["mine" if "دليل نية البحث" in (kw.get("document", "") + prompt) else "comp"]
    usage = {"input_tokens": 10, "output_tokens": 5, "total_tokens": 15, "cost_usd": 0.0001, "stop_reason": "tool_use"}
    return (json.dumps(payload, ensure_ascii=False), usage) if return_usage else json.dumps(payload)


def _events(body: str) -> dict:
    events = {}
    for chunk in body.split("\n\n"):
        lines = chunk.strip().splitlines()
        if len(lines) >= 2 and lines[0].startswith("event: "):
            events.setdefault(lines[0][7:], []).append(json.loads(lines[1][6:]))
    return events


def test_full_analysis_reports_blocked_sources_clusters_and_excel(monkeypatch):
    monkeypatch.setattr(analyze.config, "ANTHROPIC_API_KEY", "sk-ant-server")
    monkeypatch.setattr(analyze, "_url_result", _fake_url_result)
    monkeypatch.setattr(claude, "call_claude_api", _fake_ai)
    monkeypatch.setattr(verification, "call_claude_api", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("offline")))
    from app import app

    client = app.test_client()
    response = client.post("/analyze", data={
        "model": "default", "language": "ar", "ui_language": "ar",
        "my_url": "https://mine.example/a",
        "comp_urls": ["https://c1.example/a", "https://c2.example/b", "https://c3.example/c", "https://blocked.example/x"],
    })
    events = _events(response.get_data(as_text=True))
    assert "error" not in events, events.get("error")
    done = events["done"][0]

    report = {item["url"]: item["status"] for item in done["extraction_report"]}
    assert report["https://blocked.example/x"] == "blocked"
    assert any("محجوب" in label for label in done["quality"]["failed_sources"])
    assert len(done["competitor_breakdowns"]) == 3, "blocked page must be excluded, not analysed"

    clusters = done["semantic_clusters"]
    technical = next(c for c in clusters if c["pillar_entity"] == "السيو التقني")
    assert technical["is_gap"] and technical["status"] == "missing" and technical["competitor_count"] == 3
    assert {e["name"] for e in technical["supporting_entities"]} == {"الفهرسة", "خريطة الموقع"}
    gap_clusters = {gap["cluster"] for gap in done["semantic_gaps"]}
    assert technical["cluster_name"] in gap_clusters, "entity gaps carry the unified cluster name"

    entity = next(e for e in done["master_entities"] if e["name"] == "نية البحث")
    assert entity["grounded"] and entity["mentions"] >= 1 and entity["placements"]

    token = done["excel_file"].rsplit("/", 1)[-1]
    xlsx = client.get(f"/download/{token}")
    assert xlsx.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx.data))
    for sheet in ("All Entities", "العناقيد الدلالية", "فجوات العناقيد", "الفجوات الدلالية", "حالة قراءة المصادر"):
        assert sheet in workbook.sheetnames
    rows = list(workbook["فجوات العناقيد"].iter_rows(min_row=4, values_only=True))
    assert any(row[1] == "السيو التقني" for row in rows)


def test_blocked_my_site_fails_with_a_clear_message(monkeypatch):
    monkeypatch.setattr(analyze.config, "ANTHROPIC_API_KEY", "sk-ant-server")
    monkeypatch.setattr(analyze, "_url_result", _fake_url_result)
    from app import app
    response = app.test_client().post("/analyze", data={
        "model": "default", "language": "ar", "ui_language": "ar",
        "my_url": "https://blocked.example/mine", "comp_urls": ["https://c1.example/a"]})
    events = _events(response.get_data(as_text=True))
    assert "محجوب" in events["error"][0]["error"]


def test_preview_shows_status_per_source(monkeypatch):
    monkeypatch.setattr(analyze, "_url_result", _fake_url_result)
    from app import app
    response = app.test_client().post("/preview", data={
        "my_url": "https://mine.example/a", "comp_urls": ["https://blocked.example/x"], "ui_language": "ar"})
    sources = response.get_json()["sources"]
    assert sources[0]["status"] == "ok" and "text" in sources[0]
    assert sources[1]["status"] == "blocked" and "محجوب" in sources[1]["error"]


def test_api_key_comes_only_from_env(monkeypatch):
    """A key sent by the browser is ignored; without a .env key the server refuses."""
    for name in ("ANTHROPIC_API_KEY", "MISTRAL_API_KEY", "GEMINI_API_KEY"):
        monkeypatch.setattr(analyze.config, name, "", raising=False)
    from app import app
    client = app.test_client()
    for route in ("/analyze", "/analyze-serp"):
        response = client.post(route, data={"api_key": "sk-ant-from-browser", "model": "default",
                                             "ui_language": "ar", "comp_urls": ["https://c1.example/a"]})
        assert response.status_code == 400
        assert ".env" in response.get_json()["error"]
